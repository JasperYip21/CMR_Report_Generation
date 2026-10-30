"""Classify supported measurements in the CMR non-stress Word template."""

import argparse
import re
from pathlib import Path
from typing import Any

from docx import Document

try:
    from AI_agent.clinical_pdf_parser import parse_clinical_report_pdf
    from AI_agent.reference_ranges import (
        atrial_max_volume_interval,
        atrial_reference_interval,
        reference_profile,
    )
except ModuleNotFoundError:
    from clinical_pdf_parser import parse_clinical_report_pdf
    from reference_ranges import (
        atrial_max_volume_interval,
        atrial_reference_interval,
        reference_profile,
    )


_HERE = Path(__file__).resolve().parent
DEFAULT_INPUT = _HERE / "AI_test_2026-09-22_Clinical_Report_2026-09-23.pdf"
DEFAULT_TEMPLATE = _HERE / "attachments" / "CMR Non-Stress Template.docx"
DEFAULT_OUTPUT = _HERE / "CMR_Non_Stress_Report_Final.docx"


def _number(value: Any) -> float | None:
    if isinstance(value, dict):
        value = value.get("value")
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _classify(value: Any, reference_range: tuple[float, float]) -> str | None:
    number = _number(value)
    if number is None:
        return None
    low, high = reference_range
    return "normal" if low <= number <= high else "abnormal"


def _combine_classifications(*values: str | None) -> str | None:
    if not values or any(value is None for value in values):
        return None
    return "abnormal" if "abnormal" in values else "normal"


def _classify_metric_set(
    measurements: dict[str, Any], ranges: dict[str, tuple[float, float]]
) -> str | None:
    if not ranges or any(_number(measurements.get(key)) is None for key in ranges):
        return None
    return _combine_classifications(
        *(_classify(measurements[key], value) for key, value in ranges.items())
    )


def _classify_mass(data: dict[str, Any], profile: dict[str, Any]) -> str | None:
    mass_ranges = {}
    if profile["lv_mass"] is not None:
        mass_ranges["myomass_diast"] = profile["lv_mass"]
    if profile["lv_mass_bsa"] is not None:
        mass_ranges["myomass_bsa"] = profile["lv_mass_bsa"]
    return _classify_metric_set(data.get("lv", {}), mass_ranges)


def _classify_atrium(
    data: dict[str, Any],
    chamber: str,
    sex: str | None,
    age: int | None,
) -> str | None:
    measurements = data.get(chamber, {})
    statuses = []
    volume_fields = (
        (("max_volume_ml", False), ("max_volume_bsa", True))
        if chamber == "la"
        else ()
    )
    for metric, indexed in volume_fields:
        interval = atrial_max_volume_interval(sex, age, indexed=indexed)
        if interval is not None and _number(measurements.get(metric)) is not None:
            statuses.append(_classify(measurements[metric], interval))

    view = measurements.get("area_view")
    if view:
        for metric, indexed in (("area_cm2", False), ("area_bsa", True)):
            interval = atrial_reference_interval(
                chamber=chamber,
                view=view,
                sex=sex,
                age=age,
                indexed=indexed,
            )
            if interval is not None and _number(measurements.get(metric)) is not None:
                statuses.append(_classify(measurements[metric], interval))
    return _combine_classifications(*statuses) if statuses else None


def classify_report_measurements(data: dict[str, Any]) -> dict[str, list[str | None]]:
    """Classify values using the matched supplied sex/age/population profile."""
    profile = reference_profile(data)
    if profile is None:
        return {}
    lv_function = _classify_metric_set(data.get("lv", {}), profile["lv"])
    rv_function = _classify_metric_set(data.get("rv", {}), profile["rv"])
    lv_mass = _classify_mass(data, profile)
    study = data.get("study_info", {})
    la_size = _classify_atrium(
        data, "la", study.get("sex"), study.get("age")
    )
    ra_size = _classify_atrium(
        data, "ra", study.get("sex"), study.get("age")
    )
    return {
        "Left Atrium": [la_size],
        "Right Atrium": [ra_size],
        "Left Ventricle": [
            lv_function,
            lv_mass,
            None,
            lv_mass,
        ],
        "Right Ventricle": [rv_function],
        "Conclusion": [_combine_classifications(lv_function, rv_function), None],
    }


def _replace_classified_placeholders(
    paragraph: Any, values: list[str | None]
) -> None:
    """Replace classified [] tokens only; preserve other text and formatting."""
    token_index = 0

    def replace_match(match: re.Match[str]) -> str:
        nonlocal token_index
        replacement = values[token_index] if token_index < len(values) else None
        token_index += 1
        return replacement if replacement is not None else match.group()

    for run in paragraph.runs:
        run.text = re.sub(r"\[\]", replace_match, run.text)


def _write_report(
    data: dict[str, Any], template_path: str | Path, output_filename: str | Path
) -> Path:
    classifications = classify_report_measurements(data)
    document = Document(template_path)

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Left Atrium:"):
            values = classifications.get("Left Atrium")
        elif text.startswith("Right Atrium:"):
            values = classifications.get("Right Atrium")
        elif text.startswith("Left Ventricle:"):
            values = classifications.get("Left Ventricle")
        elif text.startswith("Right Ventricle:"):
            values = classifications.get("Right Ventricle")
        elif text == "[] LV & RV volumes and systolic function.":
            values = classifications.get("Conclusion")
        else:
            values = None
        if values is not None:
            _replace_classified_placeholders(paragraph, values)

    output = Path(output_filename)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def generate_non_stress_report(
    pdf_file_path: str | Path = DEFAULT_INPUT,
    template_path: str | Path = DEFAULT_TEMPLATE,
    output_filename: str | Path = DEFAULT_OUTPUT,
) -> Path:
    """Classify supported measurements and save the minimally changed template."""
    data = parse_clinical_report_pdf(pdf_file_path)
    return _write_report(data, template_path, output_filename)


def _unclassified_fields(
    classifications: dict[str, list[str | None]]
) -> list[str]:
    if not classifications:
        return [
            "all classification placeholders (no matching sex/age reference profile)"
        ]
    fields = []
    for section, label, index in (
        ("Left Atrium", "LA size", 0),
        ("Right Atrium", "RA size", 0),
        ("Left Ventricle", "LV volumes and systolic function", 0),
        ("Left Ventricle", "LV hypertrophy", 1),
        ("Left Ventricle", "LV wall thickness", 2),
        ("Left Ventricle", "LV mass", 3),
        ("Right Ventricle", "RV volumes and systolic function", 0),
        ("Conclusion", "combined LV/RV function", 0),
        ("Conclusion", "scar/infarct", 1),
    ):
        values = classifications.get(section, [])
        if index >= len(values) or values[index] is None:
            fields.append(label)
    return fields


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Classify supported CMR measurements in the non-stress template."
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="clinical report PDF")
    parser.add_argument(
        "--template",
        type=Path,
        default=DEFAULT_TEMPLATE,
        help="CMR non-stress DOCX template",
    )
    parser.add_argument(
        "--output", type=Path, default=DEFAULT_OUTPUT, help="Generated DOCX path"
    )
    args = parser.parse_args()
    data = parse_clinical_report_pdf(args.input)
    result = _write_report(data, args.template, args.output)
    print(f"Generated report: {result}")
    unclassified = _unclassified_fields(classify_report_measurements(data))
    if unclassified:
        print(
            "Template placeholders left unchanged: "
            + ", ".join(unclassified)
            + "."
        )


if __name__ == "__main__":
    main()
