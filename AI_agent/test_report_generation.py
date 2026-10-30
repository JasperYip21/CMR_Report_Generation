import tempfile
import unittest
from pathlib import Path

from docx import Document

from AI_agent.autofill import (
    DEFAULT_INPUT,
    DEFAULT_TEMPLATE,
    classify_report_measurements,
    generate_non_stress_report,
)
from AI_agent.clinical_pdf_parser import parse_clinical_report_pdf, parse_clinical_report_text
from AI_agent.cvi42_parser import parse_cvi42_report


class Cvi42ParserTests(unittest.TestCase):
    def test_reads_lv_and_rv_metrics_and_metadata_without_cross_assignment(self):
        source = "\n".join(
            (
                "Patient\tExample",
                "Clinical Results LV",
                "EF\t61.4\t%",
                "EDV\t148.8\tml",
                "EDV/BSA\t98.2\tml/m2",
                "MyoMass_diast\t92.5\tg",
                "SAX 3D LV peak wall thickness",
                "all slices, diastolic phase:",
                "slice\tphase\tpeak wall thickness (mm)",
                "total peak wall thickness (all slices, diastolic phase):",
                "5\t25\t11.75",
                "Clinical Results RV",
                "EF\t63.2\t%",
                "EDV\t146.0\tml",
            )
        )

        parsed = parse_cvi42_report(source)

        self.assertEqual(parsed["study_info"]["patient"], "Example")
        self.assertEqual(parsed["lv"]["ejection_fraction_pct"], 61.4)
        self.assertEqual(parsed["lv"]["edv"]["value"], 148.8)
        self.assertEqual(parsed["lv"]["edv_bsa"], 98.2)
        self.assertEqual(parsed["lv"]["myomass_diast"], 92.5)
        self.assertEqual(parsed["lv"]["max_wall_thickness"], 11.75)
        self.assertEqual(parsed["rv"]["ejection_fraction_pct"], 63.2)
        self.assertEqual(parsed["rv"]["edv"]["value"], 146.0)

    def test_reads_real_export_measurements(self):
        source_path = Path(__file__).with_name(
            "AI_test_2026-09-22_Scientific_Report_2026-09-23.txt"
        )
        parsed = parse_cvi42_report(source_path.read_bytes().decode("utf-16"))

        self.assertAlmostEqual(parsed["lv"]["ejection_fraction_pct"], 61.358777)
        self.assertAlmostEqual(parsed["lv"]["edv"]["value"], 148.835882)
        self.assertAlmostEqual(parsed["lv"]["edv_bsa"], 98.228489)
        self.assertAlmostEqual(parsed["lv"]["max_wall_thickness"], 11.749346)
        self.assertAlmostEqual(parsed["lv"]["myomass_diast"], 92.479882)
        self.assertAlmostEqual(parsed["rv"]["ejection_fraction_pct"], 63.233172)
        self.assertAlmostEqual(parsed["rv"]["edv"]["value"], 146.034656)

    def test_extracts_supported_fields_from_supplied_clinical_pdf(self):
        parsed = parse_clinical_report_pdf(DEFAULT_INPUT)

        self.assertEqual(parsed["study_info"]["sex"], "Female")
        self.assertEqual(parsed["study_info"]["age"], 82)
        self.assertEqual(parsed["study_info"]["height_cm"], 145)
        self.assertEqual(parsed["study_info"]["weight_kg"], 57.0)
        self.assertEqual(parsed["study_info"]["bsa_m2"], 1.52)
        self.assertAlmostEqual(parsed["aorta"]["sinus_line_measurement_mm"], 26.7)
        self.assertAlmostEqual(parsed["pulmonary_artery"]["main_line_measurement_mm"], 33.23)
        self.assertAlmostEqual(parsed["la"]["contour_area_cm2"], 31.477)
        self.assertAlmostEqual(parsed["ra"]["contour_area_cm2"], 17.078)
        self.assertAlmostEqual(parsed["lv"]["ejection_fraction_pct"], 61.359)
        self.assertAlmostEqual(parsed["lv"]["edv_bsa"], 98.228)
        self.assertAlmostEqual(parsed["lv"]["myomass_bsa"], 61.035)
        self.assertAlmostEqual(parsed["lv"]["max_wall_thickness"], 11.438)
        self.assertAlmostEqual(parsed["rv"]["ejection_fraction_pct"], 63.233)
        self.assertAlmostEqual(parsed["rv"]["edv"]["value"], 146.035)
        self.assertAlmostEqual(parsed["rv"]["edv_bsa"], 96.380)

    def test_text_parser_leaves_missing_values_unset(self):
        parsed = parse_clinical_report_text(
            "Patient Sex: F\nAge: 82 Y.\nHeight: 145 cm\n"
            "Weight: 57.0 kg\nBSA: 1.52 m2\n"
        )

        self.assertEqual(parsed["study_info"]["sex"], "Female")
        self.assertEqual(parsed["study_info"]["age"], 82)
        self.assertEqual(parsed["study_info"]["bsa_m2"], 1.52)
        self.assertEqual(parsed["lv"], {})
        self.assertEqual(parsed["rv"], {})


class ReferenceRangeTests(unittest.TestCase):
    def test_classifies_supported_sample_metrics_against_female_60_plus_ranges(self):
        parsed = parse_clinical_report_pdf(DEFAULT_INPUT)
        classifications = classify_report_measurements(parsed)

        self.assertEqual(
            classifications,
            {
                "Left Ventricle": ["abnormal", "normal", None, "normal"],
                "Right Ventricle": ["abnormal"],
                "Conclusion": ["abnormal", None],
            },
        )

    def test_does_not_classify_when_sex_or_age_reference_does_not_match(self):
        data = {
            "study_info": {"sex": "Female", "age": 59},
            "lv": {
                "edv": {"value": 148.8},
                "esv": {"value": 40},
                "sv": {"value": 90},
                "ejection_fraction_pct": 60,
                "myomass_diast": 100,
            },
            "rv": {
                "edv": {"value": 100},
                "esv": {"value": 30},
                "sv": {"value": 70},
                "ejection_fraction_pct": 60,
            },
        }

        self.assertEqual(classify_report_measurements(data), {})

        data["study_info"] = {"sex": "Male", "age": 82}
        self.assertEqual(classify_report_measurements(data), {})

    def test_requires_all_ventricle_measurements_for_combined_classification(self):
        data = {
            "study_info": {"sex": "Female", "age": 82},
            "lv": {
                "edv": {"value": 148.8},
                "esv": {"value": 40},
                "sv": {"value": 90},
                "ejection_fraction_pct": 60,
                "myomass_diast": 100,
            },
            "rv": {
                "edv": {"value": 100},
                "esv": {"value": 30},
                "sv": {"value": 70},
            },
        }

        classifications = classify_report_measurements(data)

        self.assertEqual(classifications["Left Ventricle"], ["normal", "normal", None, "normal"])
        self.assertEqual(classifications["Right Ventricle"], [None])
        self.assertEqual(classifications["Conclusion"], [None, None])


class ReportGenerationTests(unittest.TestCase):
    def test_fills_template_from_pdf_and_preserves_template_structure(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "report.docx"
            generated = generate_non_stress_report(
                DEFAULT_INPUT, DEFAULT_TEMPLATE, output_path
            )

            template = Document(DEFAULT_TEMPLATE)
            output = Document(generated)
            self.assertEqual(len(output.paragraphs), len(template.paragraphs))
            self.assertEqual(len(output.tables), len(template.tables))
            for index in (3, 7):
                self.assertEqual(
                    output.paragraphs[index].runs[0].bold,
                    template.paragraphs[index].runs[0].bold,
                )
            output_text = "\n".join(p.text for p in output.paragraphs)
            self.assertIn("abnormal LV volumes and systolic function.", output_text)
            self.assertIn("normal LV hypertrophy.", output_text)
            self.assertIn("Maximum LV wall thickness is [] mm", output_text)
            self.assertIn("normal LV mass.", output_text)
            self.assertIn("abnormal RV volumes and systolic function.", output_text)
            self.assertIn("abnormal LV & RV volumes and systolic function.", output_text)
            self.assertIn("[] LGE or infarct", output_text)
            self.assertEqual(
                output.paragraphs[22].text.strip(),
                "Left Ventricle: abnormal LV volumes and systolic function. "
                "normal LV hypertrophy. Maximum LV wall thickness is [] mm in "
                "end-diastole. normal LV mass.",
            )
            self.assertEqual(
                output.paragraphs[24].text,
                "Right Ventricle: abnormal RV volumes and systolic function.",
            )
            self.assertEqual(
                output.paragraphs[43].text,
                "abnormal LV & RV volumes and systolic function.",
            )
            for index, (original, generated_paragraph) in enumerate(
                zip(template.paragraphs, output.paragraphs)
            ):
                if index not in {22, 24, 43}:
                    self.assertEqual(generated_paragraph.text, original.text)
            self.assertEqual(
                [[cell.text for cell in row.cells] for row in output.tables[0].rows],
                [[cell.text for cell in row.cells] for row in template.tables[0].rows],
            )


if __name__ == "__main__":
    unittest.main()
